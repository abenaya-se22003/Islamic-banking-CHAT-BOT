"""
Step 8: Report Generation Tool
===============================
Provides the generate_report tool called by Claude via function calling/tool use.
Generates a formatted Microsoft Word document (.docx) based on verified
document excerpts and saves it to backend/generated_reports/.
"""

import os
import re
from datetime import datetime
from typing import Dict, Any

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ---------------------------------------------------------------------------
# Directories setup
# ---------------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPORTS_DIR = os.path.join(BASE_DIR, "generated_reports")
os.makedirs(REPORTS_DIR, exist_ok=True)

SHARIAH_DISCLAIMER = (
    "Disclaimer: This document is generated for informational purposes only based on "
    "approved Shariah-compliant banking documentation. It does not constitute a formal "
    "financial contract, offer, or binding commitment. Shariah interpretations and product terms "
    "may be updated periodically. Please consult with a qualified Shariah advisor or authorized "
    "bank officer before entering into any financial transaction."
)


def _sanitize_filename(name: str) -> str:
    """
    Sanitizes topic string to be safe for filenames.
    """
    clean = re.sub(r"[^\w\s-]", "", name).strip()
    clean = re.sub(r"[-\s]+", "_", clean)
    return clean[:50] or "report"


def generate_report(topic: str, content_summary: str) -> Dict[str, Any]:
    """
    Creates a formatted Word (.docx) report with:
    - Topic title
    - Generation timestamp
    - Formatted content summary
    - Fixed Shariah compliance disclaimer

    Args:
        topic: The title/subject of the report.
        content_summary: Detailed text content extracted strictly from verified documents.

    Returns:
        dict with status, filename, and download_url (or error message on failure).
    """
    try:
        doc = Document()

        # Set default page margins (1 inch)
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # 1. Main Title
        title_para = doc.add_paragraph()
        title_run = title_para.add_run(topic)
        title_run.font.name = "Calibri"
        title_run.font.size = Pt(22)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(16, 52, 96)  # Deep Islamic banking navy/blue
        title_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        title_para.paragraph_format.space_after = Pt(4)

        # Subtitle / Metadata
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        sub_para = doc.add_paragraph()
        sub_run = sub_para.add_run(f"Islamic Banking Report • Generated on {now_str}")
        sub_run.font.name = "Calibri"
        sub_run.font.size = Pt(10)
        sub_run.font.italic = True
        sub_run.font.color.rgb = RGBColor(100, 100, 100)
        sub_para.paragraph_format.space_after = Pt(18)

        # Divider line
        divider_para = doc.add_paragraph()
        divider_run = divider_para.add_run("―" * 45)
        divider_run.font.color.rgb = RGBColor(200, 200, 200)
        divider_para.paragraph_format.space_after = Pt(12)

        # 2. Section: Summary
        heading_summary = doc.add_paragraph()
        h_run = heading_summary.add_run("Summary & Findings")
        h_run.font.name = "Calibri"
        h_run.font.size = Pt(14)
        h_run.font.bold = True
        h_run.font.color.rgb = RGBColor(16, 52, 96)
        heading_summary.paragraph_format.space_after = Pt(8)

        # Content paragraphs
        for paragraph_text in content_summary.strip().split("\n"):
            clean_para = paragraph_text.strip()
            if not clean_para:
                continue

            p = doc.add_paragraph()
            # Handle bullet points if present
            if clean_para.startswith(("-", "•", "*")):
                p.paragraph_format.left_indent = Inches(0.25)
                run = p.add_run("• " + clean_para.lstrip("-•* "))
            else:
                run = p.add_run(clean_para)

            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(30, 30, 30)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(6)

        # 3. Fixed Shariah Disclaimer Section
        doc.add_paragraph().paragraph_format.space_after = Pt(14)
        disc_head = doc.add_paragraph()
        disc_head_run = disc_head.add_run("Shariah Advisory & Legal Disclaimer")
        disc_head_run.font.name = "Calibri"
        disc_head_run.font.size = Pt(11)
        disc_head_run.font.bold = True
        disc_head_run.font.color.rgb = RGBColor(120, 80, 20)  # Amber/Bronze tone
        disc_head.paragraph_format.space_after = Pt(4)

        disc_para = doc.add_paragraph()
        disc_run = disc_para.add_run(SHARIAH_DISCLAIMER)
        disc_run.font.name = "Calibri"
        disc_run.font.size = Pt(9.5)
        disc_run.font.italic = True
        disc_run.font.color.rgb = RGBColor(110, 110, 110)
        disc_para.paragraph_format.space_after = Pt(12)

        # Save the document
        timestamp_slug = datetime.now().strftime("%Y%m%d_%H%M%S")
        sanitized_topic = _sanitize_filename(topic)
        filename = f"{sanitized_topic}_{timestamp_slug}.docx"
        file_path = os.path.join(REPORTS_DIR, filename)

        doc.save(file_path)

        download_url = f"/reports/{filename}"
        return {
            "status": "success",
            "filename": filename,
            "download_url": download_url
        }

    except Exception as exc:
        return {
            "status": "error",
            "message": f"Failed to generate Word report: {str(exc)}"
        }
